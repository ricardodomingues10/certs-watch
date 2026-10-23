"""CMDB final: lógica do workflow n8n de reconciliação, ficheiros gerados, publicação e comparação."""
from dataclasses import replace
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from zoneinfo import ZoneInfo

from certwatch import cmdb_final as cf
from certwatch.db import session_scope
from certwatch.models import CmdbRecord, FinalDecision, FinalInput, FinalRun

TZ = ZoneInfo("Europe/Lisbon")
TODAY = date(2026, 10, 21)
P = "[Digital Certificate] : "


def ms(d: date, hour: int = 12) -> int:
    return int(datetime(d.year, d.month, d.day, hour, tzinfo=TZ).timestamp() * 1000)


def ci(uid, cn, valid_to, issuer="Corp CA01", serial="AA01", **extra):
    props = {"display_label": cn or extra.pop("display_label", ""), "name": extra.pop("name", cn), "common_name": cn,
             "valid_to": ms(valid_to) if valid_to else None, "issuer": issuer, "issuer_name": issuer,
             "serial_number": serial, "certificate_type": "SSL", "localizacao": "IIS", "aplicacao": "Portal",
             "responsabilidade_instalacao": "DSTI", "created_by": "admin", "description": "", **extra}
    return {"ucmdbId": f"u{uid}", "globalId": f"g{uid}", "type": "digital_certificate", "properties": props}


CIS = [
    ci(1, "igual.exemplo.pt", date(2027, 3, 1)),
    ci(2, "renovado.exemplo.pt", date(2026, 11, 10), serial="0A01"),
    ci(3, "www.exemplo.pt", date(2027, 6, 1), issuer="Antiga CA"),          # regressão -> servidor
    ci(4, "*.exemplo.pt", date(2027, 6, 1)),                                # regressão wildcard -> mantém
    ci(5, "mudo.exemplo.pt", date(2027, 6, 1)),                             # regressão sem resposta -> mantém
    ci(6, "", date(2026, 11, 1), name="antigo.exemplo.pt"),                 # sem CommonName, encontrado pelo Name
    ci(7, "duplo.exemplo.pt", date(2026, 1, 5), serial="D1"),
    ci(8, "duplo.exemplo.pt", date(2026, 11, 20), serial="D2"),
]


def ca_csv(rows):
    head = '"Req ID:","Requester Name:","Issued Common Name:","Certificate Expiration Date:","Serial Number:","Request Disposition:"'
    return ("\n".join([head] + [",".join(f'"{x}"' for x in r) for r in rows]) + "\n").encode("utf-8")


CA1 = ca_csv([
    ("10", "DOM\\a", "igual.exemplo.pt", "01/03/2027 10:00", "AA01", "20 -- Issued"),
    ("11", "DOM\\a", "renovado.exemplo.pt", "10/11/2027 09:00", "0a:02", "20 -- Issued"),
    ("12", "DOM\\a", "www.exemplo.pt", "01/02/2027", "BB", "20 -- Issued"),
    ("13", "DOM\\a", "*.exemplo.pt", "01/02/2027", "CC", "20 -- Issued"),
    ("14", "DOM\\a", "mudo.exemplo.pt", "01/02/2027", "DD", "20 -- Issued"),
    ("15", "DOM\\a", "antigo.exemplo.pt", "05/12/2027", "EE", "20 -- Issued"),
    ("16", "DOM\\a", "duplo.exemplo.pt", "20/12/2027", "D3", "20 -- Issued"),
    ("17", "DOM\\a", "novo-ca1.exemplo.pt", "01/01/2028", "N1", "20 -- Issued"),
    ("18", "DOM\\a", "revogado.exemplo.pt", "01/01/2028", "R1", "21 -- Revoked"),
    ("19", "DOM\\a", "expirado.exemplo.pt", "01/01/2025", "E1", "20 -- Issued"),
    ("20", "DOM\\a", "semdata.exemplo.pt", "", "S1", "20 -- Issued"),
    ("21", "DOM\\a", "SERVIDOR01", "01/01/2028", "S2", "20 -- Issued"),
    ("22", "DOM\\a", "partido.exemplo.pt", "31/31/2027", "S3", "20 -- Issued"),
])


def ca2_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Req ID:", "Issued Common Name:", "Certificate Expiration Date:", "Serial Number:", "Request Disposition:"])
    ws.append([30, "novo-ca2.exemplo.pt", datetime(2028, 2, 2, 8, 0), "N2", "20 -- Issued"])
    ws.append([31, "serial-excel.exemplo.pt", 46818.75, "N3", "20 -- Issued"])   # 2028-03-06, a hora é ignorada
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


AVI = (
    "Nome,Subject_CommonName,Validity_NotAfter,SerialNumber,Issuer_CommonName,SubjectAlternativeName\n"
    "vs-novo,novo-avi.exemplo.pt,2028-04-10 12:00:00,N4,DigiCert Global G2,novo-avi.exemplo.pt\n"
    "vs-ca1,novo-ca1.exemplo.pt,2028-01-01T00:00:00Z,N1,Corp CA01,\n"
    "vs-sem-cn,,2028-05-05,N5,DigiCert Global G2,\n"
).encode("utf-8")


@pytest.fixture()
def settings(env, tmp_path):
    imp = tmp_path / "import"
    imp.mkdir()
    return replace(env, import_dir=imp, ucmdb_min_cis=1,
                   import_issuers={"CA1": "Corp CA01", "CA2": "Corp CA02"})


def live(host):
    if host == "www.exemplo.pt":
        return {"status": "OK", "not_after": datetime(2027, 9, 30, 23, 0, tzinfo=timezone.utc),
                "serial": "FF01", "issuer_cn": "Nova CA", "origem": "teste"}
    return {"status": "ERROR", "detail": "timeout", "origem": "teste"}


def build(settings, cis=CIS):
    import copy
    table = cf.build_cmdb_table(copy.deepcopy(cis), settings)
    sources = {
        "CA1": cf.read_source("CA1", "ca1.csv", CA1, TODAY)[0],
        "CA2": cf.read_source("CA2", "ca2.xlsx", ca2_xlsx(), TODAY)[0],
        "AVI": cf.read_source("AVI", "avi.csv", AVI, TODAY)[0],
    }
    return cf.reconcile_final(table, sources, settings, live, TODAY)


def by_cn(res, cn):
    H = res.table.h
    return [(i, r) for i, r in enumerate(res.table.rows) if r[H("CommonName")] == cn]


def test_parse_day_never_swaps():
    assert cf.parse_day("10/04/2028").day == date(2028, 4, 10)
    assert cf.parse_day("10-04-2028 00:00:00").day == date(2028, 4, 10)
    assert cf.parse_day("13/25/2028").status == "INVALIDO"
    assert cf.parse_day("20280410").day == date(2028, 4, 10)
    assert cf.parse_day(46818.75).day == date(2028, 3, 6)
    assert cf.parse_day("2028-04-10T23:30:00Z").day == date(2028, 4, 10)
    assert cf.parse_day("").status == "VAZIO"
    assert cf.parse_day("amanhã").status == "NAO_RECONHECIDO"
    assert cf.parse_day("01/01/1990").status == "ANO_SUSPEITO"


def test_columns_from_api(settings):
    t = cf.build_cmdb_table([dict(c) for c in CIS], settings)
    assert t.headers[:3] == [P + "Display Label", P + "ValidTo", P + "Aplicação"]
    assert P + "Localização" in t.headers and P + "Responsabilidade de instalação" in t.headers
    assert not any("created" in h.lower() for h in t.headers)            # Created By fica de fora
    r = t.rows[0]
    assert r[P + "ValidTo"] == "01-03-2027 12:00:00"                      # epoch ms -> hora de Lisboa
    assert r[P + "Global Id"] == "g1" and r[P + "Aplicação"] == "Portal"
    assert "Subject Alternative Names" in t.diag["colunas_vazias"]


def test_column_overrides_file(settings):
    (settings.import_dir / "cmdb-colunas.csv").write_text(
        "atributo;coluna\naplicacao;\nresponsabilidade_instalacao;Responsabilidade de instalação\nx_custom;Campo X\n",
        encoding="utf-8")
    cis = [ci(1, "a.exemplo.pt", date(2027, 1, 1), x_custom="v")]
    t = cf.build_cmdb_table(cis, settings)
    assert t.rows[0][P + "Campo X"] == "v"
    assert t.rows[0][P + "Aplicação"] == ""                               # excluída no ficheiro
    assert t.headers[0] == P + "Responsabilidade de instalação"            # ordem do ficheiro


def test_reconcile_like_workflow(settings):
    res = build(settings)
    H = res.table.h
    # igual: nada muda
    (i, r), = by_cn(res, "igual.exemplo.pt")
    assert i not in res.mods
    # renovado: data e série novas, Description auditável
    (i, r), = by_cn(res, "renovado.exemplo.pt")
    assert r[H("ValidTo")] == "10-11-2027 00:00:00" and r[H("SerialNumber")] == "0a:02"
    assert r[H("Description")] == ("Changed with n8n - Data de expiracao alterado de 10/11/2026 para 10/11/2027; "
                                   "SerialNumber alterado de 0A01 para 0a:02")
    # regressão validada no servidor: o servidor manda (data, série, emissor)
    (i, r), = by_cn(res, "www.exemplo.pt")
    assert r[H("ValidTo")] == "30-09-2027 00:00:00" and r[H("SerialNumber")] == "FF01"
    assert r[H("Issuer")] == r[H("IssuerName")] == "Nova CA"
    assert r[H("Description")].startswith("Changed with n8n - Validado via TLS; Data de expiracao alterado de 01/06/2027 para 30/09/2027")
    # wildcard e servidor sem resposta: a CMDB fica como estava
    for cn in ("*.exemplo.pt", "mudo.exemplo.pt"):
        (i, r), = by_cn(res, cn)
        assert i not in res.mods and r[H("ValidTo")] == "01-06-2027 12:00:00"
    decisions = {x["cn"]: x["decisao"] for x in res.live}
    assert decisions == {"www.exemplo.pt": "ATUALIZAR_COM_SERVIDOR", "*.exemplo.pt": "MANTER_CMDB_WILDCARD",
                         "mudo.exemplo.pt": "MANTER_CMDB_SEM_RESPOSTA"}
    # sem CommonName na CMDB: encontrado pelo Name, o CommonName é preenchido na MESMA linha
    (i, r), = by_cn(res, "antigo.exemplo.pt")
    assert res.table.ids[i] == "u6" and "CommonName preenchido com antigo.exemplo.pt" in r[H("Description")]
    # duas linhas com o mesmo nome: muda a de data mais próxima
    rows = dict(by_cn(res, "duplo.exemplo.pt"))
    changed = [i for i in rows if i in res.mods]
    assert len(changed) == 1 and res.table.ids[changed[0]] == "u8"
    # novos: CA1/CA2 com o emissor configurado, AVI com o Issuer_CommonName; CA1 e AVI com o mesmo certificado = 1 linha
    new = {res.table.rows[i][H("CommonName")]: res.table.rows[i] for i in res.new_rows}
    assert set(new) == {"novo-ca1.exemplo.pt", "novo-ca2.exemplo.pt", "serial-excel.exemplo.pt", "novo-avi.exemplo.pt"}
    assert new["novo-ca1.exemplo.pt"][H("IssuerName")] == "Corp CA01"
    assert new["novo-ca2.exemplo.pt"][H("Issuer")] == "Corp CA02"
    assert new["serial-excel.exemplo.pt"][H("ValidTo")] == "06-03-2028 00:00:00"
    a = new["novo-avi.exemplo.pt"]
    assert (a[H("IssuerName")], a[H("Global Id")], a[H("Localização")], a[H("Certificate Status")],
            a[H("Certificate Type")], a[H("Description")]) == (
        "DigiCert Global G2", "x", "AVI", "Ativo", "SSL", "Changed with n8n - novo certificado")
    assert a[H("Subject Alternative Names")] == "novo-avi.exemplo.pt"
    # relatório de erros
    errs = {(p["fonte"], p["cn"]): p["erro"] for p in res.problems}
    assert errs[("CA1", "revogado.exemplo.pt")].startswith("Pedido não emitido")
    assert errs[("CA1", "expirado.exemplo.pt")] == "Certificado já expirado"
    assert errs[("CA1", "semdata.exemplo.pt")] == "Data de expiração vazia"
    assert errs[("CA1", "SERVIDOR01")].startswith("Nome sem domínio")
    assert errs[("CA1", "partido.exemplo.pt")].startswith("Data inválida")
    assert errs[("AVI", "vs-sem-cn")].startswith("Nome sem domínio")     # sem CN usa o Nome, como no workflow
    s = res.summary
    assert (s["novos"], s["alterados"], s["regressoes"], s["regressoes_atualizadas"]) == (4, 4, 3, 1)


def test_outputs_and_validation(settings):
    res = build(settings)
    delta = cf._delta_rows(res)
    assert len(delta) == 8 and not cf.validate_delta(delta)
    csv = cf.build_csv(delta, True)
    text = csv.decode("utf-8")
    assert not csv.startswith(b"\xef\xbb\xbf") and text.endswith("\r\n") and text.count("\r\n") == 9
    assert text.split("\r\n")[0].startswith('"Display Label"|"ValidTo"|"Aplicação"|')
    assert text.split("\r\n")[0].count("|") == 25
    sem = cf.build_csv(delta, False).decode("utf-8")
    assert sem.count("\r\n") == 8 and '"' not in sem and sem.split("\r\n")[0].count("|") == 25
    # aspas, pipes e quebras de linha dentro dos valores
    row = dict(delta[0], Description='a "b" | c\nd')
    assert '"a ""b""  /  c d"' in cf.build_csv([row], True).decode()     # | -> " / ", como no workflow
    assert "|a b  /  c d|" in cf.build_csv([row], False).decode()
    # novo sem número de série -> CSV recusados
    bad = [dict(delta[0], **{"Global Id": "x", "SerialNumber": ""})]
    assert "sem Issuer, IssuerName ou SerialNumber" in cf.validate_delta(bad)[0]
    dup = [dict(delta[0]), dict(delta[0])]
    assert "Chave primária repetida" in cf.validate_delta(dup)[0]
    # Excel: sem linhas extra, cores por estado, texto nunca vira fórmula
    res.table.rows[0][res.table.h("Description")] = "=1+1"
    x = cf.build_xlsx(res)
    ws = load_workbook(BytesIO(x)).worksheets[0]
    assert ws.title == "CMDB" and ws.max_row == len(res.table.rows) + 1 and ws["A1"].value == P + "Display Label"
    back = {r["cn"]: r for r in cf.read_final_xlsx(x)}
    assert back["novo-avi.exemplo.pt"]["cor"] == "novo" and back["renovado.exemplo.pt"]["cor"] == "alterado"
    assert back["igual.exemplo.pt"]["cor"] is None
    desc_col = res.table.headers.index(res.table.h("Description")) + 1
    assert ws.cell(row=2, column=desc_col).data_type == "s"


def test_expired_row_is_red(settings):
    cis = [ci(1, "velho.exemplo.pt", date(2026, 1, 1))]
    t = cf.build_cmdb_table(cis, settings)
    res = cf.reconcile_final(t, {"CA1": [], "CA2": [], "AVI": []}, settings, live, TODAY)
    assert res.expired_rows == {0}
    assert cf.read_final_xlsx(cf.build_xlsx(res))[0]["cor"] == "expirado"


def test_empty_cmdb_stops(settings):
    t = cf.build_cmdb_table([], settings)
    with pytest.raises(cf.FinalError):
        cf.reconcile_final(t, {}, settings, live, TODAY)


def test_generate(settings, monkeypatch):
    (settings.import_dir / "ca1_export.csv").write_bytes(CA1)
    (settings.import_dir / "avi_export.csv").write_bytes(AVI)
    with session_scope() as s:
        s.add(FinalInput(source="CA2", name="ca2.xlsx", sha256="x", size=1, data=ca2_xlsx(), uploaded_by="admin"))
        inputs = cf.current_inputs(s, settings)
        assert {k: v["origem"] for k, v in inputs.items()} == {"CA1": "diretoria", "AVI": "diretoria", "CA2": "carregado"}
        run_id = cf.start_run(s, "admin").id
        with pytest.raises(cf.FinalError):
            cf.start_run(s, "admin")                       # só uma geração de cada vez

    def fetch(_settings):
        import copy
        cis = copy.deepcopy(CIS)
        table = cf.build_cmdb_table(cis, _settings)
        recs = [{"ucmdb_id": c["ucmdbId"], "global_id": c["globalId"], "common_name": c["properties"]["common_name"],
                 "valid_to": datetime.fromtimestamp(c["properties"]["valid_to"] / 1000, tz=timezone.utc).isoformat(),
                 "issuer": c["properties"]["issuer"], "issuer_name": c["properties"]["issuer_name"]} for c in cis]
        return table, recs

    cf.generate(run_id, settings, fetch=fetch, live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run_id)
        assert run.status == "pronto", run.error
        assert run.summary["novos"] == 4 and run.csv and run.csv_sem and run.xlsx
        assert run.inputs["CA2"]["origem"] == "carregado"
        assert s.scalar(select(CmdbRecord).where(CmdbRecord.common_name == "renovado.exemplo.pt")) is not None



def test_generate_blocked_and_error(settings):
    (settings.import_dir / "ca1.csv").write_bytes(CA1)
    (settings.import_dir / "avi.csv").write_bytes(AVI)
    with session_scope() as s:
        run_id = cf.start_run(s, "admin").id
    # falta CA2 -> erro claro
    cf.generate(run_id, settings, fetch=lambda st: (cf.build_cmdb_table([dict(c) for c in CIS], st), []),
                live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run_id)
        assert run.status == "erro" and "CA2" in run.error
    # emissor de CA2 não configurado e não encontrado noutra fonte -> "Por verificar" (os CSV são gerados)
    (settings.import_dir / "ca2.xlsx").write_bytes(ca2_xlsx())
    st2 = replace(settings, import_issuers={"CA1": "Corp CA01"})
    with session_scope() as s:
        run_id = cf.start_run(s, "admin").id
    import copy
    cf.generate(run_id, st2, fetch=lambda st: (cf.build_cmdb_table(copy.deepcopy(CIS), st), []),
                live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run_id)
        assert run.status == "pronto" and run.csv and not run.validation
        novo = next(c for c in run.changes if c["cn"] == "novo-ca2.exemplo.pt")
        assert novo["issuer_name"] == "Por verificar" and "IssuerName" in novo["pendente"]
        assert run.summary["por_verificar"] >= 1
        line = next(x for x in run.csv.decode().splitlines() if "novo-ca2.exemplo.pt" in x)
        assert '"Por verificar"' in line


def test_web_flow(settings, monkeypatch):
    import copy

    from fastapi.testclient import TestClient

    from certwatch.web import app as webapp

    monkeypatch.setattr(webapp, "get_settings", lambda: settings)

    def fake_bg(run_id, st, only=None):
        cf.generate(run_id, st, fetch=lambda s2: (cf.build_cmdb_table(copy.deepcopy(CIS), s2), []),
                    live_factory=lambda s, s2: live, today=TODAY, only=only)
    monkeypatch.setattr(cf, "generate_in_background", fake_bg)
    AUTH = ("admin", "teste-admin-123")
    (settings.import_dir / "ca1.csv").write_bytes(CA1)
    with TestClient(webapp.app) as c:
        r = c.get("/cmdb-final", auth=AUTH)
        assert r.status_code == 200 and "Falta CA2, AVI" in r.text
        token = c.cookies["cw_csrf"]
        # carregar fontes pela página
        r = c.post("/cmdb-final/fonte/CA2", data={"csrf": token}, files={"ficheiro": ("ca2.xlsx", ca2_xlsx())}, auth=AUTH)
        assert r.status_code == 200 and "ca2.xlsx" in r.text
        r = c.post("/cmdb-final/fonte/AVI", data={"csrf": token}, files={"ficheiro": ("avi.exe", b"MZ")}, auth=AUTH)
        assert "Formato não aceite" in r.text
        r = c.post("/cmdb-final/fonte/AVI", data={"csrf": "x"}, files={"ficheiro": ("avi.csv", AVI)}, auth=AUTH)
        assert r.status_code == 403
        c.post("/cmdb-final/fonte/AVI", data={"csrf": token}, files={"ficheiro": ("avi.csv", AVI)}, auth=AUTH)
        # gerar
        r = c.post("/cmdb-final/gerar", data={"csrf": token}, auth=AUTH, follow_redirects=False)
        assert r.status_code == 303
        url = r.headers["location"]
        r = c.get(url, auth=AUTH)
        assert "Pronta" in r.text and "O que foi comparado" in r.text and "Aplicar na UCMDB" in r.text
        assert "<b>8</b> alterações" in r.text and "Quando aparece aqui" in r.text
        for v in ("novos", "alterados", "servidor", "confirmar", "so_na_cmdb", "azure", "erros", "detalhes",
                  "alteracoes", "colunas"):
            assert c.get(f"{url}?vista={v}", auth=AUTH).status_code == 200
        assert "MANTER" not in c.get(f"{url}?vista=servidor", auth=AUTH).text   # texto legível, não o código
        novos = c.get(f"{url}?vista=novos", auth=AUTH).text
        assert "novo-avi.exemplo.pt" in novos and "Não existe na CMDB" in novos and "renovado.exemplo.pt" not in novos
        alt = c.get(f"{url}?vista=alterados", auth=AUTH).text
        assert "renovado.exemplo.pt" in alt and "Na CMDB está 10-11-2026" in alt
        assert "Sem exportação do Azure" in c.get(f"{url}?vista=azure", auth=AUTH).text
        r = c.get(f"{url}/ficheiro/csv", auth=AUTH)
        assert r.content.startswith(b'"Display Label"|')
        assert c.get(f"{url}/ficheiro/xlsx", auth=AUTH).content[:2] == b"PK"
        for what in ("alteracoes", "erros", "servidor", "colunas"):
            assert c.get(f"{url}/exportar/{what}", auth=AUTH).status_code == 200
        # excluir uma alteração e voltar a incluí-la
        assert 'name="incluir"' in novos
        r = c.post(f"{url}/selecao", data={"csrf": token, "k": ["1", "2"], "incluir": ["2"],
                                           "next": f"{url}?vista=alterados"}, auth=AUTH)
        assert r.status_code == 200 and "1 excluída no total" in r.text and "<b>7</b> alterações" in r.text
        r = c.post(f"{url}/selecao", data={"csrf": token, "k": ["1"], "incluir": ["1"]}, auth=AUTH)
        assert "<b>8</b> alterações" in r.text
        assert c.post(f"{url}/selecao", data={"csrf": token, "k": ["x"]}, auth=AUTH).status_code == 400
        # páginas antigas
        assert c.get("/cmdb-final/comparar", auth=AUTH).status_code in (404, 422)
        page = c.get("/cmdb-final", auth=AUTH).text
        assert "Reconciliar agora" in page and "Renovações" in page and "Azure" in page


def _ucmdb(topology_query_status=200):
    import json as _json

    import httpx
    full = {c["ucmdbId"]: c["properties"] for c in CIS}
    calls = {"tq": 0, "ci": 0}

    def handler(request: httpx.Request):
        path = request.url.path
        if path.endswith("/authenticate"):
            return httpx.Response(200, json={"token": "t"})
        if path.endswith("/topology"):
            # layout do TQL só com os campos confiáveis (como o TQL do certwatch)
            cis = [{"ucmdbId": c["ucmdbId"], "globalId": c["globalId"], "type": c["type"],
                    "properties": {k: c["properties"][k] for k in ("common_name", "valid_to", "issuer", "issuer_name")}}
                   for c in CIS]
            return httpx.Response(200, json={"cis": cis, "relations": []})
        if "/dataModel/ci/" in path:
            calls["ci"] += 1
            uid = path.rsplit("/", 1)[1]
            return httpx.Response(200, json={"ucmdbId": uid, "properties": {**full[uid], "root_class": "digital_certificate",
                                                                             "data_note": "x", "password_hint": "nunca"}})
        if path.endswith("/topologyQuery"):
            calls["tq"] += 1
            if topology_query_status != 200:
                return httpx.Response(topology_query_status, json={"error": "sem permissão"})
            layout = _json.loads(request.content)["nodes"][0]["layout"]
            return httpx.Response(200, json={"cis": [{"ucmdbId": u, "properties": {k: p.get(k) for k in layout}}
                                                     for u, p in full.items()], "relations": []})
        return httpx.Response(404)
    return httpx.Client(base_url="https://cmdb.test/rest-api", transport=httpx.MockTransport(handler)), calls


def test_fetch_cmdb_completes_attributes(settings):
    client, calls = _ucmdb()
    table, recs = cf.fetch_cmdb_table(settings, client)
    assert table.diag["metodo"] == "topologyQuery" and calls == {"tq": 1, "ci": 1}
    r = table.rows[0]
    assert r[P + "Aplicação"] == "Portal" and r[P + "Responsabilidade de instalação"] == "DSTI"
    assert r[P + "Global Id"] == "g1" and r[P + "ValidTo"] == "01-03-2027 12:00:00"
    assert not any("root" in h or "data_note" in h or "password" in h for h in table.headers)
    assert recs[0]["common_name"] == "igual.exemplo.pt" and len(recs) == len(CIS)
    # sem permissão para a consulta ad-hoc: lê CI a CI
    client, calls = _ucmdb(403)
    table, _ = cf.fetch_cmdb_table(settings, client)
    assert table.diag["metodo"] == "ci" and calls["ci"] == 1 + len(CIS)
    assert table.rows[1][P + "Aplicação"] == "Portal"
    assert any("topologyQuery" in a for a in table.diag["avisos"])


def _ct(session, name, not_after, triage="por_triar", serial="C1", issuer="R11", names=None):
    from certwatch.models import CtCert
    session.add(CtCert(primary_name=name, names=names or [name], issuer_cn=issuer, serial_hex=serial,
                       not_before=not_after - (date(2027, 1, 1) - date(2026, 10, 1)),
                       not_after=not_after, triage=triage, domain="exemplo.pt"))


def _fetch(st):
    import copy
    return cf.build_cmdb_table(copy.deepcopy(CIS), st), []


def test_internet_source_and_selection(settings):
    (settings.import_dir / "ca1.csv").write_bytes(CA1)
    (settings.import_dir / "ca2.xlsx").write_bytes(ca2_xlsx())
    (settings.import_dir / "avi.csv").write_bytes(AVI)
    utc = timezone.utc
    with session_scope() as s:
        _ct(s, "ct-novo.exemplo.pt", datetime(2027, 8, 1, 10, tzinfo=utc), "adicionar", "ab01", names=["ct-novo.exemplo.pt", "alt.exemplo.pt"])
        _ct(s, "ct-triar.exemplo.pt", datetime(2027, 8, 2, tzinfo=utc), "por_triar", "ab02")
        _ct(s, "ct-triar.exemplo.pt", datetime(2027, 2, 2, tzinfo=utc), "por_triar", "ab00")      # mais antigo: ignorado
        _ct(s, "ct-ignorado.exemplo.pt", datetime(2027, 8, 3, tzinfo=utc), "ignorar", "ab03")
        _ct(s, "velho.exemplo.pt", datetime(2025, 8, 3, tzinfo=utc), "adicionar", "ab04")          # expirado
        _ct(s, "igual.exemplo.pt", datetime(2027, 12, 1, tzinfo=utc), "por_triar", "ab05")        # renovação de um que existe
        items, info = cf.ct_source(s, settings, TODAY)
        assert sorted(i.cn for i in items) == ["ct-novo.exemplo.pt", "ct-triar.exemplo.pt", "igual.exemplo.pt"]
        assert info["triagem"] == {"adicionar": 1, "por_triar": 2, "investigar": 0} and info["ignorados_na_triagem"] == 1
        assert next(i for i in items if i.cn == "ct-triar.exemplo.pt").serial == "AB02"
        run_id = cf.start_run(s, "admin").id
    cf.generate(run_id, settings, fetch=_fetch, live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run_id)
        assert run.status == "pronto", run.error
        ch = {c["cn"]: c for c in run.changes}
        assert ch["ct-novo.exemplo.pt"]["incluir"] and ch["ct-novo.exemplo.pt"]["fonte"] == "INTERNET"
        assert not ch["ct-triar.exemplo.pt"]["incluir"] and "Adicionar à CMDB" in ch["ct-triar.exemplo.pt"]["porque"]
        assert ch["igual.exemplo.pt"]["tipo"] == "alterado" and not ch["igual.exemplo.pt"]["incluir"]
        assert "ct-ignorado.exemplo.pt" not in ch and run.inputs["INTERNET"]["validas"] == 3
        csv = run.csv.decode()
        assert "ct-novo.exemplo.pt" in csv and "ct-triar.exemplo.pt" not in csv and "igual.exemplo.pt" not in csv
        assert '"R11"|"R11"|"AB01"' in csv and '"INTERNET"' in csv                         # emissor do certificado
        assert run.summary["novos"] == 5 and run.summary["novos_total"] == 6 and run.summary["excluidos"] == 2
        final = {r["cn"]: r for r in cf.read_final_xlsx(run.xlsx)}
        assert "ct-triar.exemplo.pt" not in final and final["igual.exemplo.pt"]["cor"] is None
        assert final["igual.exemplo.pt"]["validto"] == "01-03-2027 12:00:00"                   # ficou como na CMDB

        # incluir o "por triar" e excluir a renovação de renovado.exemplo.pt
        keys = {c["cn"]: c["chave"] for c in run.changes}
        shown = [keys["ct-triar.exemplo.pt"], keys["renovado.exemplo.pt"]]
        assert cf.apply_selection(s, run, shown, {keys["ct-triar.exemplo.pt"]}, "ana") == 2
        assert cf.apply_selection(s, run, shown, {keys["ct-triar.exemplo.pt"]}, "ana") == 0       # nada mudou
        csv = run.csv.decode()
        assert "ct-triar.exemplo.pt" in csv and "renovado.exemplo.pt" not in csv
        final = {r["cn"]: r for r in cf.read_final_xlsx(run.xlsx)}
        assert final["renovado.exemplo.pt"]["cor"] is None and final["renovado.exemplo.pt"]["validto"] == "10-11-2026 12:00:00"
        assert run.summary["alterados"] == 3 and run.summary["excluidos"] == 2

    # a escolha fica para a geração seguinte
    with session_scope() as s:
        run2 = cf.start_run(s, "admin").id
    cf.generate(run2, settings, fetch=_fetch, live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run2)
        ch = {c["cn"]: c for c in run.changes}
        assert ch["ct-triar.exemplo.pt"]["incluir"] and "incluída por ana" in ch["ct-triar.exemplo.pt"]["porque"]
        assert not ch["renovado.exemplo.pt"]["incluir"]
        old = s.get(FinalRun, run_id)
        with pytest.raises(cf.FinalError):                      # só a mais recente pode ser alterada
            cf.apply_selection(s, old, [0], set(), "ana")


def test_exclusion_updates_pending(settings):
    (settings.import_dir / "ca1.csv").write_bytes(CA1)
    (settings.import_dir / "ca2.xlsx").write_bytes(ca2_xlsx())
    (settings.import_dir / "avi.csv").write_bytes(AVI)
    st = replace(settings, import_issuers={"CA1": "Corp CA01"}, final_ct=False)    # CA2 sem emissor -> "Por verificar"
    with session_scope() as s:
        run_id = cf.start_run(s, "admin").id
    cf.generate(run_id, st, fetch=_fetch, live_factory=lambda s, st: live, today=TODAY)
    with session_scope() as s:
        run = s.get(FinalRun, run_id)
        assert run.status == "pronto" and run.summary["por_verificar"] >= 1
        bad = [c["chave"] for c in run.changes if c["fonte"] == "CA2"]
        cf.apply_selection(s, run, bad, set(), "ana")
        assert run.status == "pronto" and run.csv and run.summary["por_verificar"] == 0
        assert "Por verificar" not in run.csv.decode()


def test_missing_serial_and_issuer(settings):
    """Linhas novas (Global Id = x): Issuer e IssuerName sempre preenchidos ("Por verificar" se não houver em nenhuma
    fonte); o SerialNumber nunca é inventado. Antes disso, procura noutras fontes e nos certificados já vistos."""
    cis = [ci(1, "renovado.exemplo.pt", date(2026, 11, 10), serial="0A01"),
           ci(2, "igual.exemplo.pt", date(2027, 3, 1), serial="AA01"),
           ci(3, "carregado-antes.exemplo.pt", date(2027, 9, 1), serial="x7")]      # já na CMDB de outra vez
    t = cf.build_cmdb_table(cis, settings)
    st = replace(settings, import_issuers={})
    S = cf.SourceItem
    sources = {
        "CA1": [S("CA1", 2, "renovado.exemplo.pt", date(2027, 11, 10)),          # renovado, sem série -> x<n>
                S("CA1", 3, "igual.exemplo.pt", date(2027, 3, 1)),               # mesmo certificado -> nada muda
                S("CA1", 4, "novo1.exemplo.pt", date(2028, 1, 1), serial="N1"),  # sem emissor em lado nenhum
                S("CA1", 5, "novo2.exemplo.pt", date(2028, 2, 1)),               # outra fonte tem os dados
                S("CA1", 6, "novo3.exemplo.pt", date(2028, 3, 1)),               # certwatch já o viu
                S("CA1", 7, "novo4.exemplo.pt", date(2028, 4, 1))],              # sem série nem emissor
        "AVI": [S("AVI", 2, "novo2.exemplo.pt", date(2028, 2, 1), serial="N2", issuer="DigiCert G2")],
        "INTERNET": [S("INTERNET", 1, "net1.exemplo.pt", date(2028, 5, 1), triage="adicionar"),       # sem série nem emissor
                     S("INTERNET", 2, "net2.exemplo.pt", date(2028, 6, 1), triage="adicionar")],
    }
    known = lambda cn, day: {"issuer": "Corp CA03", "serial": "C3"} if (cn, day) == ("novo3.exemplo.pt", date(2028, 3, 1)) else {}
    res = cf.reconcile_final(t, sources, st, live, TODAY, known)
    by = {c["cn"]: c for c in res.changes}
    ren = by["renovado.exemplo.pt"]
    # outro certificado (outra data) sem série: a antiga já não serve -> x8 (a seguir ao x7 da CMDB)
    assert ren["serial"] == "x8" and ren["pendente"] == ["SerialNumber"] and ren["global_id"] != "x"
    assert "SerialNumber alterado de 0A01 para x8" in ren["descricao"]
    assert "igual.exemplo.pt" not in by
    assert (by["novo1.exemplo.pt"]["issuer"], by["novo1.exemplo.pt"]["issuer_name"], by["novo1.exemplo.pt"]["serial"]) == \
        ("Por verificar", "Por verificar", "N1")
    assert by["novo1.exemplo.pt"]["pendente"] == ["IssuerName"] and by["novo1.exemplo.pt"]["global_id"] == "x"
    assert (by["novo2.exemplo.pt"]["issuer_name"], by["novo2.exemplo.pt"]["serial"], by["novo2.exemplo.pt"]["pendente"]) == \
        ("DigiCert G2", "N2", [])
    assert (by["novo3.exemplo.pt"]["issuer_name"], by["novo3.exemplo.pt"]["serial"]) == ("Corp CA03", "C3")
    assert by["novo4.exemplo.pt"]["serial"] == "" and by["novo4.exemplo.pt"]["pendente"] == ["IssuerName", "SerialNumber"]
    assert not any(cf.is_placeholder_serial(c["serial"]) for c in res.changes if c["fonte"] != "INTERNET" and c["tipo"] == "novo")
    # sem série na internet: x9, x10, cada um com a sua chave
    for n, x in (("net1.exemplo.pt", "x9"), ("net2.exemplo.pt", "x10")):
        assert (by[n]["issuer_name"], by[n]["serial"], by[n]["global_id"]) == ("Por verificar", x, "x")
        assert by[n]["pendente"] == ["IssuerName", "SerialNumber"]
    # sem série: fica fora dos CSV por omissão; o resto é gerado
    for c in res.changes:
        c["incluir"], c["porque"] = cf.default_include(c, {})
    assert not by["novo4.exemplo.pt"]["incluir"] and "SerialNumber" in by["novo4.exemplo.pt"]["porque"]
    assert by["net1.exemplo.pt"]["incluir"]
    xlsx, csv, csv_sem, problems, counts = cf.build_outputs(cf.unpack_state(cf.pack_state(res)), res.changes)
    assert not problems and csv and counts["por_verificar"] == 4
    assert counts["por_verificar_repetidos"] == 0        # x9 e x10: chaves diferentes
    net = next(ln for ln in csv.decode().splitlines() if "net1.exemplo.pt" in ln)
    assert '"Por verificar"|"Por verificar"|"x9"' in net
    text = csv.decode()
    assert "novo4.exemplo.pt" not in text
    line = next(ln for ln in text.splitlines() if "novo1.exemplo.pt" in ln).split("|")
    head = text.splitlines()[0].split("|")
    col = {h.strip('"'): n for n, h in enumerate(head)}
    assert line[col["Issuer"]] == line[col["IssuerName"]] == '"Por verificar"' and line[col["SerialNumber"]] == '"N1"'
    assert line[col["Global Id"]] == '"x"'
    # incluído à força sem série -> bloqueia (chave primária incompleta)
    by["novo4.exemplo.pt"]["incluir"] = True
    assert cf.build_outputs(cf.unpack_state(cf.pack_state(res)), res.changes)[3]


def test_csv_format_is_fixed(settings):
    res = build(settings)
    _, csv, csv_sem, problems, _ = cf.build_outputs(cf.unpack_state(cf.pack_state(res)), res.changes)
    expected = ["Display Label", "ValidTo", "Aplicação", "Localização", "Responsabilidade de instalação",
                "Certificate Status", "Certificado Gerido", "Ambiente", "Name", "CommonName", "Issuer", "IssuerName",
                "SerialNumber", "Certificate Type", "IsSelfSigned", "Global Id", "Link para a Documentação",
                "Subject Alternative Names", "Organization", "Version", "Type", "Subject", "CreateOn", "Hostname",
                "Description", "OrganizationUnit"]
    lines = csv.decode("utf-8").split("\r\n")
    assert [h.strip('"') for h in lines[0].split("|")] == expected
    assert all(len(ln.split("|")) == 26 for ln in lines if ln)
    assert all(len(ln.split("|")) == 26 for ln in csv_sem.decode("utf-8").split("\r\n") if ln)
    # mesmo sem alterações, o ficheiro com cabeçalho tem o cabeçalho
    assert cf.build_csv([], True).decode().split("\r\n")[0].count("|") == 25


def test_known_lookup_from_certwatch(settings):
    from certwatch.models import Certificate, CtCert
    with session_scope() as s:
        s.add(Certificate(sha256="a" * 64, subject_cn="visto.exemplo.pt", issuer_cn="Corp CA01", serial_hex="0abc",
                          not_after=datetime(2028, 3, 1, 23, 30, tzinfo=timezone.utc), sans_dns=["alt.exemplo.pt"]))
        s.add(CtCert(names=["ct.exemplo.pt"], primary_name="ct.exemplo.pt", issuer_cn="R11", serial_hex="dd",
                     not_after=datetime(2028, 4, 1, 12, 0, tzinfo=timezone.utc)))
    with session_scope() as s:
        look = cf.make_known_lookup(s, settings)
        assert look("visto.exemplo.pt", date(2028, 3, 2)) == {"issuer": "Corp CA01", "serial": "0ABC"}   # 23:30 UTC = dia 2 em Lisboa? ±1 dia
        assert look("alt.exemplo.pt", date(2028, 3, 1))["serial"] == "0ABC"
        assert look("ct.exemplo.pt", date(2028, 4, 1)) == {"issuer": "R11", "serial": "DD"}
        assert look("visto.exemplo.pt", date(2028, 6, 1)) == {}
        assert look("outro.exemplo.pt", date(2028, 3, 1)) == {}


def test_run_page_shows_pending(settings, monkeypatch):
    from fastapi.testclient import TestClient
    from certwatch.web import app as webapp
    st = replace(settings, import_issuers={"CA1": "Corp CA01"}, final_ct=False)
    monkeypatch.setattr(webapp, "get_settings", lambda: st)
    (st.import_dir / "ca1.csv").write_bytes(CA1)
    (st.import_dir / "ca2.xlsx").write_bytes(ca2_xlsx())
    (st.import_dir / "avi.csv").write_bytes(AVI)
    with session_scope() as s:
        run_id = cf.start_run(s, "admin").id
    cf.generate(run_id, st, fetch=_fetch, live_factory=lambda s, x: live, today=TODAY)
    with TestClient(webapp.app) as c:
        page = c.get(f"/cmdb-final/{run_id}?vista=confirmar", auth=("admin", "teste-admin-123")).text
        assert "Por confirmar" in page and "Os ficheiros CSV não foram gerados" not in page
        assert "novo-ca2.exemplo.pt" in page and "novo-avi.exemplo.pt" not in page


def test_duplicate_key_cases(settings):
    """O caso "Chave primária repetida": dois CIs com a mesma série renovados sem série -> x<n> (sem repetição);
    duas alterações com a mesma série real -> ficam de fora por omissão, com a indicação de com quem colidem."""
    cis = [ci(1, "a.exemplo.pt", date(2026, 11, 10), serial="586D2F04", issuer="YR1"),
           ci(2, "b.exemplo.pt", date(2026, 11, 10), serial="586D2F04", issuer="YR1"),
           ci(3, "c.exemplo.pt", date(2026, 12, 1), serial="AB", issuer="Corp CA01"),
           ci(4, "d.exemplo.pt", date(2026, 12, 1), serial="CD", issuer="Corp CA01")]
    t = cf.build_cmdb_table(cis, settings)
    S = cf.SourceItem
    sources = {"INTERNET": [S("INTERNET", 1, "a.exemplo.pt", date(2027, 11, 10), issuer="YR1", triage="adicionar"),
                            S("INTERNET", 2, "b.exemplo.pt", date(2027, 11, 10), issuer="YR1", triage="adicionar")],
               "AVI": [S("AVI", 2, "c.exemplo.pt", date(2027, 12, 1), serial="EF", issuer="Corp CA01"),
                       S("AVI", 3, "d.exemplo.pt", date(2027, 12, 1), serial="EF", issuer="Corp CA01")]}
    res = cf.reconcile_final(t, sources, settings, live, TODAY)
    by = {c["cn"]: c for c in res.changes}
    assert {by["a.exemplo.pt"]["serial"], by["b.exemplo.pt"]["serial"]} == {"x1", "x2"}
    for c in res.changes:
        c["incluir"], c["porque"] = cf.default_include(c, {})
    cf.flag_duplicate_keys(res.changes, {})
    # fica uma (a de um CI que já existe / nome mais curto); a outra sai, a dizer com quem colide
    assert by["c.exemplo.pt"]["incluir"] and "d.exemplo.pt" in by["c.exemplo.pt"]["chave_repetida"]
    assert not by["d.exemplo.pt"]["incluir"] and "mesmo certificado que c.exemplo.pt" in by["d.exemplo.pt"]["porque"]
    assert by["a.exemplo.pt"]["incluir"] and "chave_repetida" not in by["a.exemplo.pt"]
    xlsx, csv, csv_sem, problems, counts = cf.build_outputs(cf.unpack_state(cf.pack_state(res)), res.changes)
    assert not problems and csv
    # as duas incluídas à mão -> bloqueia, com os nomes e Global Ids na mensagem
    by["c.exemplo.pt"]["incluir"] = by["d.exemplo.pt"]["incluir"] = True
    problems = cf.build_outputs(cf.unpack_state(cf.pack_state(res)), res.changes)[3]
    assert problems and "c.exemplo.pt" in problems[0] and "d.exemplo.pt" in problems[0] and "Corp CA01 + EF" in problems[0]

    # na página: incluir a d troca (a c sai), e a escolha fica para a próxima reconciliação
    by["d.exemplo.pt"]["incluir"] = False
    with session_scope() as s:
        run = FinalRun(created_by="admin", status="pronto", state=cf.pack_state(res), changes=res.changes)
        s.add(run)
        s.flush()
        k = {c["cn"]: c["chave"] for c in run.changes}
        assert cf.apply_selection(s, run, [k["c.exemplo.pt"], k["d.exemplo.pt"]], {k["c.exemplo.pt"], k["d.exemplo.pt"]}, "ana")
        ch = {c["cn"]: c for c in run.changes}
        assert ch["d.exemplo.pt"]["incluir"] and not ch["c.exemplo.pt"]["incluir"] and run.status == "pronto"
        assert "mesmo certificado que d.exemplo.pt" in ch["c.exemplo.pt"]["porque"]
        decisions = {d.key: d for d in s.scalars(select(FinalDecision))}
        again = [dict(c) for c in res.changes]
        for c in again:
            c["incluir"], c["porque"] = cf.default_include(c, decisions)
        cf.flag_duplicate_keys(again, decisions)
        assert {c["cn"]: c["incluir"] for c in again if c["cn"] in ("c.exemplo.pt", "d.exemplo.pt")} == \
            {"c.exemplo.pt": False, "d.exemplo.pt": True}


def test_reconcile_one_source(settings, monkeypatch):
    """Reconciliar só uma fonte: as outras não são lidas nem obrigatórias. Sem nada carregado, usa a diretoria import/."""
    import copy

    from fastapi.testclient import TestClient

    from certwatch.web import app as webapp

    monkeypatch.setattr(webapp, "get_settings", lambda: settings)

    def fake_bg(run_id, st, only=None):
        cf.generate(run_id, st, fetch=lambda s2: (cf.build_cmdb_table(copy.deepcopy(CIS), s2), []),
                    live_factory=lambda s, s2: live, today=TODAY, only=only)
    monkeypatch.setattr(cf, "generate_in_background", fake_bg)
    AUTH = ("admin", "teste-admin-123")
    (settings.import_dir / "avi.csv").write_bytes(AVI)            # só o AVI, na diretoria
    with TestClient(webapp.app) as c:
        page = c.get("/cmdb-final", auth=AUTH).text
        assert "Reconciliar só esta fonte" in page and "import/avi.csv" in page
        token = c.cookies["cw_csrf"]
        # tudo: falta CA1 e CA2, com a indicação do nome esperado na diretoria
        r = c.post("/cmdb-final/gerar", data={"csrf": token}, auth=AUTH)
        assert "Falta o ficheiro de: CA1 (import/" in r.text
        # só o AVI
        r = c.post("/cmdb-final/gerar", data={"csrf": token, "fonte": "AVI"}, auth=AUTH, follow_redirects=False)
        url = r.headers["location"]
        page = c.get(url, auth=AUTH).text
        assert "só AVI" in page
        assert c.post("/cmdb-final/gerar", data={"csrf": token, "fonte": "XPTO"}, auth=AUTH).status_code == 400
    with session_scope() as s:
        run = s.scalar(select(FinalRun).order_by(FinalRun.id.desc()))
        assert run.status == "pronto", run.error
        assert run.summary["ambito"] == ["AVI"] and set(run.inputs) == {"AVI"}
        assert {c["fonte"] for c in run.changes} == {"AVI"}
        assert run.inputs["AVI"]["origem"] == "diretoria"
    # carregado mais recente ganha; "Voltar a usar a diretoria" esquece-o
    with TestClient(webapp.app) as c:
        c.get("/cmdb-final", auth=AUTH)
        token = c.cookies["cw_csrf"]
        c.post("/cmdb-final/fonte/AVI", data={"csrf": token}, files={"ficheiro": ("avi-novo.csv", AVI)}, auth=AUTH)
        with session_scope() as s:
            assert cf.current_inputs(s, settings)["AVI"]["origem"] == "carregado"
        page = c.get("/cmdb-final", auth=AUTH).text
        assert "Voltar a usar a diretoria import/" in page and "não usado: há um carregado mais recente" in page
        c.post("/cmdb-final/fonte/AVI/diretoria", data={"csrf": token}, auth=AUTH)
        with session_scope() as s:
            assert cf.current_inputs(s, settings)["AVI"]["origem"] == "diretoria"


def test_data_flow_azure_only():
    from types import SimpleNamespace
    run = SimpleNamespace(summary={"ambito": ["AZURE"], "cmdb_linhas": 632}, columns={"tql": "T"},
                          inputs={"AZURE": {"validas": 3, "apis_lidas": 40, "ficheiro": "azure.csv"}},
                          vistas={"azure": [{"estado": "novo"}, {"estado": "novo"}, {"estado": "igual"}]})
    flow = cf.data_flow(run)
    assert flow[0]["titulo"] == "Azure" and flow[0]["n"] == 3 and flow[1]["n"] == 40
    assert flow[2]["n"] == 3 and flow[3]["n"] == 2 and "2 não existem na CMDB" in flow[3]["det"]
