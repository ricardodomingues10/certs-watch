"""Relatório por email: intervalo e alerta à escolha, alertas e reconciliação, envio junto ou separado."""
from dataclasses import replace
from datetime import timedelta

import pytest
from cryptography.hazmat.primitives.serialization import Encoding
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO
from sqlalchemy import select

from certwatch import report as rp
from certwatch.db import session_scope, utcnow
from certwatch.jobs import _store_outcome, ingest_cmdb
from certwatch.models import EmailReport, Endpoint
from certwatch.scanner import ScanOutcome

AUTH = ("admin", "teste-admin-123")
NOW = utcnow()


class FakeSMTP:
    sent: list = []

    def __init__(self, settings):
        self.settings = settings

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def login(self, u, p):
        FakeSMTP.sent.append(("login", u))

    def send_message(self, m):
        FakeSMTP.sent.append(m)


@pytest.fixture()
def data(env, pki):
    st = replace(env, ucmdb_min_cis=1, smtp_host="smtp.exemplo.pt", smtp_from="certwatch@exemplo.pt",
                 report_allowed_domains=["exemplo.pt"], report_signature="Equipa de Segurança")
    with session_scope() as s:
        ingest_cmdb(s, st, [
            {"ucmdb_id": str(i), "global_id": str(i), "common_name": cn, "valid_to": (NOW + timedelta(days=d)).isoformat(),
             "issuer": "Corp CA01", "issuer_name": "Corp CA01"}
            for i, (cn, d) in enumerate([("app.exemplo.local", 10), ("www.exemplo.pt", 50), ("longe.exemplo.pt", 200),
                                         ("renovado.exemplo.pt", 5)])])
        # o servidor já apresenta o certificado renovado (validade muito posterior à da CMDB)
        leaf, _ = pki.leaf("renovado.exemplo.pt", sans=["renovado.exemplo.pt"], days_valid=400)
        ep = s.scalar(select(Endpoint).where(Endpoint.host == "renovado.exemplo.pt"))
        _store_outcome(s, ep, ScanOutcome("renovado.exemplo.pt", 443, chain_der=[leaf.public_bytes(Encoding.DER)],
                                          verify_ok=True, tls_version="TLSv1.3"))
    return st


def test_rows_window_alerts_and_reconciliation(data):
    with session_scope() as s:
        rows = rp.build_rows(s, data)
    by = {r["cn"]: r for r in rows}
    assert by["renovado.exemplo.pt"]["cmdb_atualizada"] == "Não" and by["renovado.exemplo.pt"]["origem"] == "Servidor"
    assert by["app.exemplo.local"]["scope"] == "interno" and by["www.exemplo.pt"]["scope"] == "externo"
    sel = rp.select_rows(rows, 0, 90)
    assert sorted(r["cn"] for r in sel["janela"]) == ["app.exemplo.local", "www.exemplo.pt"]
    assert [r["cn"] for r in sel["renovados"]] == ["renovado.exemplo.pt"]
    html = rp.email_html("alertas", sel, 0, 90, 20, data.report_signature, NOW)
    assert "ALERTA — Certificados fora do prazo de renovação (20 dias" in html and "app.exemplo.local" in html
    assert "nos próximos 90 dias" in html and "Sem pedido ativo" in html and "Equipa de Segurança" in html
    assert "renovado.exemplo.pt" not in html                                  # não é desta parte
    rec = rp.email_html("reconciliacao", sel, 0, 90, 20, "x", NOW)
    assert "1 certificado(s)" in rec and "renovado.exemplo.pt" in rec and "Excel" in rec
    assert rp.subject_for("alertas", sel, 0, 90, 20) == "ALERTA | 1 certificado(s) a 20 dias ou menos da expiração"
    assert rp.subject_for("alertas", sel, 0, 90, 5) == "Certificados a expirar nos próximos 90 dias"
    assert rp.subject_for("reconciliacao", sel, 0, 90, 20).startswith("Reconciliação de certificados | 1 renovado")
    # prazo de alerta à escolha: a 60 dias entram os dois
    assert rp.email_html("alertas", sel, 0, 90, 60, "x", NOW).count("font-weight:bold") == 2
    # retirar um certificado
    sel2 = rp.select_rows(rows, 0, 90, {"app.exemplo.local"})
    assert [r["cn"] for r in sel2["janela"]] == ["www.exemplo.pt"]
    # Excel: janela + renovados, com o estado da análise
    ws = load_workbook(BytesIO(rp.excel(sel, 20))).worksheets[0]
    vals = {r[0]: dict(zip(rp.EXCEL_COLUMNS, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(vals) == {"app.exemplo.local", "www.exemplo.pt", "renovado.exemplo.pt"}
    assert vals["renovado.exemplo.pt"]["Estado_Analise"] == "RENOVADO_NAO_REFLETIDO_CMDB"
    assert vals["app.exemplo.local"]["Estado_Analise"] == "ALERTA_20_DIAS_SEM_TICKET_ATIVO"
    assert vals["www.exemplo.pt"]["Estado_Analise"] == "SEM_TICKET"


def test_recipients(data):
    assert rp.parse_recipients("a@exemplo.pt; b@sub.exemplo.pt a@exemplo.pt", data) == ["a@exemplo.pt", "b@sub.exemplo.pt"]
    for bad in ("", "nao-e-email", "x@gmail.com", "a@exemplo.pt\r\nBcc: x@y.z"):
        with pytest.raises(rp.ReportError):
            rp.parse_recipients(bad, data)


def test_web_flow_and_send(data, monkeypatch):
    from certwatch.web import app as webapp
    monkeypatch.setattr(webapp, "get_settings", lambda: data)
    monkeypatch.setattr(webapp, "SMTP_FACTORY", FakeSMTP)
    FakeSMTP.sent = []
    with TestClient(webapp.app) as c:
        assert c.get("/relatorio", auth=AUTH).status_code == 200
        token = c.cookies["cw_csrf"]
        assert c.post("/relatorio", data={"csrf": token, "desde": "30", "ate": "10", "alerta": "20"}, auth=AUTH).status_code == 400
        r = c.post("/relatorio", data={"csrf": token, "desde": "0", "ate": "90", "alerta": "20"}, auth=AUTH, follow_redirects=False)
        url = r.headers["location"]
        page = c.get(url, auth=AUTH).text
        assert "app.exemplo.local" in page and "1 certificado</b> a 20 dias" in page and "previsao/alertas" in page
        prev = c.get(url + "/previsao/alertas", auth=AUTH)
        assert "ALERTA" in prev.text and "frame-ancestors 'self'" in prev.headers["content-security-policy"]
        assert prev.headers["x-frame-options"] == "SAMEORIGIN"
        assert c.get(url + "/excel", auth=AUTH).content[:2] == b"PK"
        # retirar um certificado
        c.post(url + "/selecao", data={"csrf": token, "k": ["app.exemplo.local", "www.exemplo.pt"], "inc": ["www.exemplo.pt"]}, auth=AUTH)
        with session_scope() as s:
            assert s.scalar(select(EmailReport)).excluded == ["app.exemplo.local"]
        assert "app.exemplo.local" not in c.get(url + "/previsao/alertas", auth=AUTH).text
        c.post(url + "/selecao", data={"csrf": token, "k": ["app.exemplo.local"], "inc": ["app.exemplo.local"]}, auth=AUTH)
        # domínio não permitido: não envia e regista a falha
        r = c.post(url + "/enviar", data={"csrf": token, "para": "x@gmail.com", "parte": ["alertas"], "confirmo": "1"}, auth=AUTH)
        assert "só são permitidos endereços de exemplo.pt" in r.text and not FakeSMTP.sent
        # as duas partes em dois emails
        r = c.post(url + "/enviar", data={"csrf": token, "para": "ricardo@exemplo.pt", "parte": ["alertas", "reconciliacao"],
                                          "modo": "separado", "confirmo": "1"}, auth=AUTH)
        assert "Enviado para ricardo@exemplo.pt" in r.text
        msgs = [m for m in FakeSMTP.sent if not isinstance(m, tuple)]
        assert len(msgs) == 2 and msgs[0]["Subject"].startswith("ALERTA |") and msgs[0]["To"] == "ricardo@exemplo.pt"
        assert not list(msgs[0].iter_attachments())
        att = list(msgs[1].iter_attachments())
        assert msgs[1]["Subject"].startswith("Reconciliação") and att[0].get_filename() == "Certificados_90_Dias.xlsx"
        # as duas juntas: um email com anexo
        FakeSMTP.sent = []
        c.post(url + "/enviar", data={"csrf": token, "para": "ricardo@exemplo.pt", "parte": ["alertas", "reconciliacao"],
                                      "modo": "junto", "confirmo": "1"}, auth=AUTH)
        msgs = [m for m in FakeSMTP.sent if not isinstance(m, tuple)]
        assert len(msgs) == 1 and list(msgs[0].iter_attachments())
        body = msgs[0].get_body(("html",)).get_content()
        assert "ALERTA" in body and "renovado.exemplo.pt" in body
        # só a reconciliação
        FakeSMTP.sent = []
        c.post(url + "/enviar", data={"csrf": token, "para": "ricardo@exemplo.pt", "parte": ["reconciliacao"], "confirmo": "1"}, auth=AUTH)
        assert len([m for m in FakeSMTP.sent if not isinstance(m, tuple)]) == 1
        page = c.get(url, auth=AUTH).text
        assert "enviado" in page and "falhou" in page
        with session_scope() as s:
            assert len(s.scalar(select(EmailReport)).sends) == 4


def test_not_configured(env):
    assert "CERTWATCH_SMTP_HOST não definido" in rp.smtp_problems(env)
    with pytest.raises(rp.ReportError):
        rp.send(env, [])


def test_internet_approval_email(data, monkeypatch):
    from certwatch.models import CtCert
    utc_now = utcnow()
    with session_scope() as s:
        def ct(name, days, triage="por_triar", seen_days_ago=2, names=None, serial="AB"):
            s.add(CtCert(primary_name=name, names=names or [name], issuer_cn="R11", serial_hex=serial,
                         not_before=utc_now - timedelta(days=10), not_after=utc_now + timedelta(days=days), triage=triage,
                         first_seen=utc_now - timedelta(days=seen_days_ago), source_ids={"crtsh": "123"}, domain="exemplo.pt"))
        ct("loja.exemplo.pt", 80, names=["loja.exemplo.pt", "www.loja.exemplo.pt"])          # nome desconhecido
        ct("www.exemplo.pt", 300, serial="CD")                                                 # nome na CMDB, outra data
        ct("www.exemplo.pt", 50, serial="EF")                                                  # igual à CMDB -> fora
        ct("ruido.exemplo.pt", 80, "ignorar")                                                   # ignorado -> fora
        ct("velho.exemplo.pt", -5)                                                               # expirado -> fora
        ct("antigo.exemplo.pt", 80, seen_days_ago=90)                                           # fora do período
    with session_scope() as s:
        rows = rp.internet_rows(s, data, 30)
        assert [(r["nome"], r["estado"]) for r in rows] == [("loja.exemplo.pt", "desconhecido"), ("www.exemplo.pt", "nome_conhecido")]
        assert len(rp.internet_rows(s, data, 0)) == 3
    assert rows[0]["link"] == "https://crt.sh/?id=123" and rows[0]["outros"] == ["www.loja.exemplo.pt"]
    sel = rp.select_rows([], 0, 90, {"ct:x"}, rows)
    html = rp.email_html(["internet"], sel, 0, 90, 20, "Equipa", NOW, 30)
    assert "2 certificado(s)" in html and "nos últimos 30 dias" in html and "Aprovado" in html
    assert html.index("loja.exemplo.pt") < html.index("Nomes que já existem na CMDB")
    assert rp.subject_for("internet", sel, 0, 90, 20) == "Aprovação | 2 certificado(s) novo(s) encontrado(s) na internet"
    ws = load_workbook(BytesIO(rp.internet_excel(sel))).worksheets[0]
    head = [c.value for c in ws[1]]
    assert "Aprovado (Sim/Não)" in head and ws.max_row == 3

    from certwatch.web import app as webapp
    monkeypatch.setattr(webapp, "get_settings", lambda: data)
    monkeypatch.setattr(webapp, "SMTP_FACTORY", FakeSMTP)
    FakeSMTP.sent = []
    with TestClient(webapp.app) as c:
        c.get("/relatorio", auth=AUTH)
        token = c.cookies["cw_csrf"]
        url = c.post("/relatorio", data={"csrf": token, "desde": "0", "ate": "90", "alerta": "20", "ct_dias": "30"},
                     auth=AUTH, follow_redirects=False).headers["location"]
        page = c.get(url, auth=AUTH).text
        assert "loja.exemplo.pt" in page and "previsao/internet" in page and "2</b> novos na internet" in page
        assert "Pedimos à equipa de Sistemas" in c.get(url + "/previsao/internet", auth=AUTH).text
        # retirar um e enviar só esta parte
        with session_scope() as s:
            ids = {r["nome"]: r["key"] for r in rp.internet_rows(s, data, 30)}
        c.post(url + "/selecao", data={"csrf": token, "k": list(ids.values()), "inc": [ids["loja.exemplo.pt"]]}, auth=AUTH)
        r = c.post(url + "/enviar", data={"csrf": token, "para": "sistemas@exemplo.pt", "parte": ["internet"], "confirmo": "1"}, auth=AUTH)
        assert "Enviado para sistemas@exemplo.pt" in r.text
        msgs = [m for m in FakeSMTP.sent if not isinstance(m, tuple)]
        assert len(msgs) == 1 and msgs[0]["Subject"] == "Aprovação | 1 certificado(s) novo(s) encontrado(s) na internet"
        assert [a.get_filename() for a in msgs[0].iter_attachments()] == ["Certificados_Internet.xlsx"]
        # as três partes num só email: dois anexos
        FakeSMTP.sent = []
        c.post(url + "/enviar", data={"csrf": token, "para": "sistemas@exemplo.pt", "parte": ["alertas", "reconciliacao", "internet"],
                                      "modo": "junto", "confirmo": "1"}, auth=AUTH)
        msgs = [m for m in FakeSMTP.sent if not isinstance(m, tuple)]
        assert len(msgs) == 1 and len(list(msgs[0].iter_attachments())) == 2
        FakeSMTP.sent = []
        c.post(url + "/enviar", data={"csrf": token, "para": "sistemas@exemplo.pt", "parte": ["alertas", "reconciliacao", "internet"],
                                      "modo": "separado", "confirmo": "1"}, auth=AUTH)
        assert len([m for m in FakeSMTP.sent if not isinstance(m, tuple)]) == 3
