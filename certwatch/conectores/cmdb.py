"""Conectores de CMDB.

Todos devolvem os certificados no formato de sempre (o da UCMDB):
    {"ucmdb_id": id estável do registo, "global_id": opcional, "common_name": CN,
     "valid_to": data ISO, "issuer": emissor, "issuer_name": nome do emissor (opcional)}
"ucmdb_id" é o identificador do registo na CMDB do cliente (o nome ficou por compatibilidade).
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import httpx

from . import Connector, ConnectorError, Field, TestResult, as_text, dig, iso

CERT_CAPS = frozenset({"certificados"})


def _record(rid, cn, valid_to, issuer, issuer_name="", global_id=None) -> dict:
    return {"ucmdb_id": str(rid or "").strip()[:64], "global_id": (str(global_id)[:64] if global_id else None),
            "common_name": as_text(cn)[:512], "valid_to": iso(valid_to), "issuer": as_text(issuer)[:512],
            "issuer_name": as_text(issuer_name)[:512]}


MAPPING_FIELDS = (
    Field("campo_id", "Campo do identificador", default="id", advanced=True,
          help="Identificador estável do registo (chave primária, sys_id...)."),
    Field("campo_cn", "Campo do nome (CN)", default="cn", advanced=True),
    Field("campo_validade", "Campo da data de expiração", default="valid_to", advanced=True,
          help="ISO, dd/mm/aaaa ou epoch: a data é reconhecida sozinha."),
    Field("campo_emissor", "Campo do emissor", default="issuer", advanced=True),
    Field("campo_emissor_nome", "Campo do nome do emissor (opcional)", default="", advanced=True),
)

MIN_FIELD = Field("min_registos", "Mínimo de registos", kind="number", default="1", advanced=True,
                  help="Abaixo disto a importação é abortada (protege de desativar certificados por engano).")


class _Mapped(Connector):
    """Conectores que leem linhas/objetos e os mapeiam por nome de campo."""

    role = "cmdb"
    capabilities = CERT_CAPS

    @property
    def min_records(self) -> int:
        try:
            return max(0, int(self.p("min_registos", "1")))
        except ValueError:
            return 1

    def map_row(self, row: dict, index: int) -> dict:
        def get(name: str):
            key = self.p(name)
            if not key:
                return None
            return dig(row, key) if "." in key else row.get(key, row.get(key.lower(), row.get(key.upper())))

        rid = get("campo_id") or f"linha-{index + 1}"
        return _record(rid, get("campo_cn"), get("campo_validade"), get("campo_emissor"), get("campo_emissor_nome") or "")

    def rows(self) -> list[dict]:
        raise NotImplementedError

    def fetch_certificates(self) -> list[dict]:
        return [self.map_row(r, i) for i, r in enumerate(self.rows())]

    def test(self) -> TestResult:
        if self.missing():
            return TestResult(False, "Falta preencher: " + ", ".join(self.missing()))
        try:
            recs = self.fetch_certificates()
        except ConnectorError as exc:
            return TestResult(False, str(exc))
        with_cn = [r for r in recs if r["common_name"]]
        with_date = [r for r in recs if r["valid_to"]]
        msg = f"{len(recs)} registos lidos · {len(with_cn)} com nome · {len(with_date)} com data de expiração"
        ok = bool(recs) and bool(with_cn)
        if recs and not with_date:
            msg += ". Nenhuma data reconhecida: confirma o campo da data."
            ok = False
        return TestResult(ok, msg, sample=recs[:8], total=len(recs))


# ---------------------------------------------------------------- UCMDB (o de sempre)
class UcmdbCmdb(Connector):
    role, kind, label = "cmdb", "ucmdb", "UCMDB (Micro Focus / OpenText)"
    description = "REST API da UCMDB com um TQL de certificados. É a configuração do .env, como até agora."
    capabilities = frozenset({"certificados", "apis", "escrita"})
    FIELDS = ()  # configurada no .env e nos segredos ucmdb_user/ucmdb_password (README, secção 1)

    @property
    def min_records(self) -> int:
        return self.settings.ucmdb_min_cis

    def configured(self) -> bool:
        return bool(self.settings.ucmdb_url and self.settings.ucmdb_tql)

    def missing(self) -> list[str]:
        return [] if self.configured() else ["CERTWATCH_UCMDB_URL e CERTWATCH_UCMDB_TQL no .env"]

    def summary(self) -> str:
        return f"{self.label} · {self.settings.ucmdb_url} · TQL {self.settings.ucmdb_tql}"

    def fetch_certificates(self) -> list[dict]:
        from ..fontes.ucmdb import fetch_certificates

        return fetch_certificates(self.settings)

    def fetch_apis(self):
        from ..fontes.ucmdb import fetch_apis

        return fetch_apis(self.settings)

    def test(self) -> TestResult:
        from ..fontes.ucmdb import UcmdbError

        try:
            recs = self.fetch_certificates()
        except (UcmdbError, httpx.HTTPError) as exc:
            return TestResult(False, str(exc))
        return TestResult(bool(recs), f"{len(recs)} certificados no TQL {self.settings.ucmdb_tql}",
                          sample=recs[:8], total=len(recs))


# ---------------------------------------------------------------- base de dados SQL
class SqlCmdb(_Mapped):
    kind, label = "sql", "Base de dados (SQL)"
    description = ("PostgreSQL, MySQL/MariaDB, SQL Server ou Oracle: uma consulta só de leitura que devolve os "
                   "certificados. Serve para uma CMDB própria ou para uma vista criada para o Atalaia.")
    FIELDS = (
        Field("url", "Ligação", required=True,
              help="postgresql+psycopg://utilizador@servidor:5432/base · mysql+pymysql://utilizador@servidor/base · "
                   "mssql+pymssql://utilizador@servidor/base · oracle+oracledb://utilizador@servidor:1521/?service_name=X"),
        Field("password", "Palavra-passe", kind="password", secret=True),
        Field("consulta", "Consulta", kind="textarea", required=True,
              default="SELECT id, cn, valid_to, issuer FROM certificados",
              help="Só SELECT (ou WITH). Os nomes das colunas são os campos abaixo."),
        *MAPPING_FIELDS,
        MIN_FIELD,
    )

    def summary(self) -> str:
        return f"{self.label} · {re.sub(r'://([^@/]*)@', '://***@', self.p('url'))}"

    def rows(self) -> list[dict]:
        from sqlalchemy import create_engine, text
        from sqlalchemy.engine import make_url
        from sqlalchemy.exc import ArgumentError, NoSuchModuleError, SQLAlchemyError

        query = self.p("consulta")
        if not re.match(r"^\s*(select|with)\b", query, re.I) or ";" in query.strip().rstrip(";"):
            raise ConnectorError("A consulta tem de ser um único SELECT (ou WITH ... SELECT).")
        try:
            url = make_url(self.p("url"))
            if self.s("password") and url.get_backend_name() != "sqlite":
                url = url.set(password=self.s("password"))
            engine = create_engine(url, pool_pre_ping=True, connect_args=_timeout_args(url.get_backend_name()))
        except (ArgumentError, NoSuchModuleError) as exc:
            raise ConnectorError(f"Ligação inválida ou controlador em falta: {exc}") from exc
        try:
            with engine.connect() as conn:
                if url.get_backend_name() == "postgresql":
                    conn.execute(text("SET TRANSACTION READ ONLY"))
                result = conn.execute(text(query.rstrip().rstrip(";")))
                cols = list(result.keys())
                return [dict(zip(cols, r, strict=False)) for r in result.fetchmany(200_000)]
        except SQLAlchemyError as exc:
            msg = str(getattr(exc, "orig", exc)).splitlines()[0][:300]
            raise ConnectorError(f"A base de dados recusou: {msg}") from exc
        finally:
            engine.dispose()


def _timeout_args(backend: str) -> dict:
    return {"postgresql": {"connect_timeout": 15}, "mysql": {"connect_timeout": 15},
            "mssql": {"login_timeout": 15}}.get(backend, {})


# ---------------------------------------------------------------- API REST / JSON genérica
AUTH_OPTIONS = (("nenhuma", "Sem autenticação"), ("bearer", "Token (Bearer)"), ("basic", "Utilizador e palavra-passe"),
                ("cabecalho", "Cabeçalho próprio (ex.: X-API-Key)"))


class RestCmdb(_Mapped):
    kind, label = "rest", "API REST / JSON"
    description = "Qualquer CMDB ou inventário com uma API que devolve JSON (GLPI, iTop, Device42, NetBox, uma API própria)."
    FIELDS = (
        Field("url", "Endereço", required=True, help="O pedido GET que devolve a lista de certificados."),
        Field("autenticacao", "Autenticação", kind="select", default="bearer", options=AUTH_OPTIONS),
        Field("utilizador", "Utilizador (autenticação básica)"),
        Field("cabecalho", "Nome do cabeçalho", default="X-API-Key", advanced=True),
        Field("token", "Token, chave ou palavra-passe", kind="password", secret=True),
        Field("lista", "Onde está a lista no JSON", default="",
              help="Caminho com pontos, ex.: result ou data.items. Vazio = a resposta é a lista."),
        Field("paginacao", "Paginação", kind="select", default="nenhuma", advanced=True,
              options=(("nenhuma", "Sem paginação"), ("offset", "offset + limit"), ("pagina", "página + tamanho"),
                       ("seguinte", "link para a página seguinte no JSON"))),
        Field("param_offset", "Parâmetro de início", default="offset", advanced=True),
        Field("param_limite", "Parâmetro de tamanho", default="limit", advanced=True),
        Field("tamanho", "Registos por página", kind="number", default="500", advanced=True),
        Field("campo_seguinte", "Caminho do link seguinte", default="next", advanced=True),
        Field("verificar_tls", "Verificar o certificado TLS", kind="bool", default="sim", advanced=True),
        *MAPPING_FIELDS,
        MIN_FIELD,
    )

    def summary(self) -> str:
        return f"{self.label} · {self.p('url')}"

    def _client(self) -> httpx.Client:
        from ..trust import ssl_context_for_api   # mesma confiança TLS (raízes internas em trust/)

        headers = {"Accept": "application/json"}
        auth = None
        mode = self.p("autenticacao", "nenhuma")
        if mode == "bearer" and self.s("token"):
            headers["Authorization"] = f"Bearer {self.s('token')}"
        elif mode == "cabecalho" and self.s("token"):
            headers[self.p("cabecalho", "X-API-Key")] = self.s("token")
        elif mode == "basic":
            auth = (self.p("utilizador"), self.s("token"))
        verify = ssl_context_for_api(self.settings) if self.p("verificar_tls", "sim") != "nao" else False
        return httpx.Client(headers=headers, auth=auth, timeout=httpx.Timeout(60.0), verify=verify,
                            follow_redirects=True)

    def get_json(self, client: httpx.Client, url: str, params: dict | None = None):
        try:
            r = client.get(url, params=params)
        except httpx.HTTPError as exc:
            raise ConnectorError(f"Sem ligação a {url}: {exc.__class__.__name__}") from exc
        if r.status_code in (401, 403):
            raise ConnectorError(f"Acesso recusado ({r.status_code}): confirma a credencial e as permissões.")
        if r.status_code != 200:
            raise ConnectorError(f"{url} respondeu {r.status_code}: {r.text[:200]}")
        try:
            return r.json()
        except ValueError as exc:
            raise ConnectorError("A resposta não é JSON.") from exc

    def _items(self, data) -> list[dict]:
        items = dig(data, self.p("lista"))
        if isinstance(items, dict):
            items = list(items.values()) if all(isinstance(v, dict) for v in items.values()) else [items]
        if not isinstance(items, list):
            raise ConnectorError(f"Não encontrei uma lista em '{self.p('lista') or '(raiz)'}' na resposta.")
        return [i for i in items if isinstance(i, dict)]

    def rows(self) -> list[dict]:
        out: list[dict] = []
        mode = self.p("paginacao", "nenhuma")
        size = int(self.p("tamanho", "500") or 500)
        with self._client() as c:
            if mode == "nenhuma":
                return self._items(self.get_json(c, self.p("url")))
            if mode == "seguinte":
                url, guard = self.p("url"), 0
                while url and guard < 2000:
                    data = self.get_json(c, url)
                    out += self._items(data)
                    nxt = dig(data, self.p("campo_seguinte", "next"))
                    url, guard = (str(nxt) if nxt else ""), guard + 1
                return out
            start = 0 if mode == "offset" else 1
            while len(out) < 200_000:
                params = {self.p("param_offset", "offset"): start, self.p("param_limite", "limit"): size}
                items = self._items(self.get_json(c, self.p("url"), params))
                out += items
                if len(items) < size:
                    break
                start += size if mode == "offset" else 1
        return out


# ---------------------------------------------------------------- ServiceNow
class ServiceNowCmdb(RestCmdb):
    kind, label = "servicenow", "ServiceNow CMDB"
    description = "Tabela de certificados do ServiceNow (Certificate Inventory, cmdb_ci_certificate) pela Table API."
    FIELDS = (
        Field("instancia", "Instância", required=True, help="https://empresa.service-now.com"),
        Field("utilizador", "Utilizador de integração", required=True),
        Field("token", "Palavra-passe", kind="password", secret=True, required=True),
        Field("tabela", "Tabela", default="cmdb_ci_certificate"),
        Field("consulta", "Filtro (sysparm_query)", default="", help="Ex.: install_status=1 (vazio = todos)."),
        Field("campo_id", "Campo do identificador", default="sys_id", advanced=True),
        Field("campo_cn", "Campo do nome (CN)", default="subject_common_name", advanced=True),
        Field("campo_validade", "Campo da data de expiração", default="valid_to", advanced=True),
        Field("campo_emissor", "Campo do emissor", default="issuer_common_name", advanced=True),
        Field("campo_emissor_nome", "Campo do nome do emissor (opcional)", default="", advanced=True),
        MIN_FIELD,
    )

    def summary(self) -> str:
        return f"{self.label} · {self.p('instancia')} · {self.p('tabela')}"

    def rows(self) -> list[dict]:
        base = self.p("instancia").rstrip("/")
        self.params.update(autenticacao="basic", verificar_tls="sim")
        fields = ",".join(x for x in (self.p("campo_id"), self.p("campo_cn"), self.p("campo_validade"),
                                       self.p("campo_emissor"), self.p("campo_emissor_nome")) if x)
        out: list[dict] = []
        offset, size = 0, 1000
        with self._client() as c:
            while len(out) < 200_000:
                params = {"sysparm_limit": size, "sysparm_offset": offset, "sysparm_fields": fields,
                          "sysparm_exclude_reference_link": "true"}
                if self.p("consulta"):
                    params["sysparm_query"] = self.p("consulta")
                data = self.get_json(c, f"{base}/api/now/table/{self.p('tabela', 'cmdb_ci_certificate')}", params)
                items = [i for i in (data.get("result") or []) if isinstance(i, dict)]
                out += items
                if len(items) < size:
                    break
                offset += size
        return out


# ---------------------------------------------------------------- ficheiro CSV / Excel
class FileCmdb(_Mapped):
    kind, label = "ficheiro", "Ficheiro CSV ou Excel"
    description = ("Uma exportação da CMDB colocada na diretoria de importação (import/), substituída sempre que "
                   "houver uma nova. Para quem não tem API ou prefere não a abrir.")
    FIELDS = (
        Field("ficheiro", "Nome do ficheiro (ou início do nome)", required=True, default="cmdb-certificados",
              help="Procura em import/ o ficheiro .csv ou .xlsx mais recente com este início."),
        Field("folha", "Folha do Excel (opcional)", advanced=True),
        *MAPPING_FIELDS,
        MIN_FIELD,
    )

    def summary(self) -> str:
        return f"{self.label} · import/{self.p('ficheiro')}*"

    def path(self) -> Path:
        d = Path(self.settings.import_dir)
        prefix = Path(self.p("ficheiro")).name.lower()
        found = sorted((f for f in d.glob("*") if f.is_file() and f.name.lower().startswith(prefix)
                        and f.suffix.lower() in (".csv", ".xlsx", ".txt")), key=lambda f: f.stat().st_mtime, reverse=True)
        if not found:
            raise ConnectorError(f"Nenhum ficheiro '{prefix}*.csv' ou '.xlsx' em {d}.")
        return found[0]

    def rows(self) -> list[dict]:
        path = self.path()
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[self.p("folha")] if self.p("folha") and self.p("folha") in wb.sheetnames else wb.active
            it = ws.iter_rows(values_only=True)
            header = [str(h or "").strip() for h in next(it, [])]
            return [dict(zip(header, r, strict=False)) for r in it if any(v not in (None, "") for v in r)]
        raw = path.read_bytes()
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
        return [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


TYPES: dict[str, type[Connector]] = {
    "ucmdb": UcmdbCmdb,
    "servicenow": ServiceNowCmdb,
    "sql": SqlCmdb,
    "rest": RestCmdb,
    "ficheiro": FileCmdb,
}
