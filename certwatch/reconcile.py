"""Leitura dos ficheiros de exportação (CA1, CA2, AVI, Azure, lista de APIs) da diretoria de importação.

Lê CSV (UTF-8, UTF-16 do certutil, Windows-1252) e Excel (.xlsx), reconhece as colunas mesmo com
acentos estragados ("LocalizaÃ§Ã£o") e interpreta datas portuguesas (DD/MM/AAAA, nunca troca dia e
mês), seriais do Excel e ISO. A comparação com a CMDB está em cmdb_final (Reconciliar e atualizar).
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from .config import Settings

ALLOWED_EXT = {".csv", ".txt", ".xlsx", ".xlsm"}
MAX_PROBLEMS = 200

COLUMNS = {
    "cn": ["issuedcommonname", "subjectcommonname", "commonname", "cn", "appregistration"],
    "cn_fallback": ["nome", "name", "displaylabel", "hostname", "enterpriseapplication"],
    "valid_to": ["certificateexpirationdate", "validitynotafter", "notafter", "validto", "expirationdate",
                 "dataexpiracao", "datadeexpiracao", "datadevalidade", "validade", "expira", "enddate"],
    "serial": ["serialnumber", "serial", "numerodeserie", "keyid"],
    # Azure AD (credenciais de App Registrations)
    "cred_type": ["credentialtype"],
    "app_id": ["appid", "applicationid", "clientid"],
    "enterprise_app": ["enterpriseapplication"],
    "start": ["startdate"],
    "status": ["status"],
    "issuer": ["issuercommonname", "issuername", "issuer", "emissor"],
    "san": ["subjectalternativename", "subjectalternativenames", "san"],
    "disposition": ["requestdisposition", "disposition"],
    "req_id": ["reqid", "requestid"],
    "requester": ["requestername"],
}

# ------------------------------------------------------------------ leitura de ficheiros
def _fix_mojibake(s: str) -> str:
    if re.search(r"Ã.|Â.|â€", s):
        try:
            fixed = s.encode("latin-1").decode("utf-8")
            if "�" not in fixed:
                return fixed
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return s


def norm_header(h) -> str:
    s = _fix_mojibake(str(h or "")).replace("﻿", "").replace(" ", " ")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"^\s*\[[^\]]*\]\s*:\s*", "", s.lower())   # "[Digital Certificate] : ValidTo"
    return re.sub(r"[^a-z0-9]", "", s)


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    if len(data) > 4 and data[1:2] == b"\x00" and data[3:4] == b"\x00":
        return data.decode("utf-16-le")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def read_table(path: Path) -> tuple[list[str], list[list]]:
    return read_table_bytes(path.read_bytes(), path.suffix)


def read_table_bytes(data: bytes, suffix: str) -> tuple[list[str], list[list]]:
    """Cabeçalho e linhas de um CSV ou Excel (o tipo vem da extensão: .xlsx/.xlsm ou texto)."""
    if suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        text = _decode(data)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = ";" if sample.count(";") > sample.count(",") else ","
        rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    # Cabeçalho = primeira linha com pelo menos duas células preenchidas
    start = next((i for i, r in enumerate(rows) if sum(1 for c in r if str(c or "").strip()) >= 2), None)
    if start is None:
        return [], []
    headers = [str(c or "").strip() for c in rows[start]]
    body = [r for r in rows[start + 1:] if any(str(c or "").strip() for c in r)]
    return headers, body


def map_columns(headers: list[str]) -> dict[str, int]:
    normed = [norm_header(h) for h in headers]
    found: dict[str, int] = {}
    for field, cands in COLUMNS.items():
        for cand in cands:
            if cand in normed:
                found[field] = normed.index(cand)
                break
    if "cn" not in found and "cn_fallback" in found:
        found["cn"] = found["cn_fallback"]
    if "valid_to" not in found:
        for i, h in enumerate(normed):
            if "expir" in h or "notafter" in h or h.endswith("validto"):
                found["valid_to"] = i
                break
    if "cn" not in found:
        for i, h in enumerate(normed):
            if "commonname" in h:
                found["cn"] = i
                break
    found.pop("cn_fallback", None)
    return found


_DMY = re.compile(r"^(\d{1,2})([/.\-])(\d{1,2})\2(\d{4})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM)?)?$", re.I)


def parse_date(value, tz: ZoneInfo) -> datetime | None:
    """Datas de exportação: datetime do Excel, serial do Excel, ISO, AAAAMMDD, DD/MM/AAAA[ HH:MM[:SS]].
    DD/MM nunca é trocado por MM/DD, exceto quando a data tem AM/PM (formato americano explícito)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return (value.replace(tzinfo=tz) if value.tzinfo is None else value).astimezone(timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=tz).astimezone(timezone.utc)
    if isinstance(value, (int, float)):
        if 20000 < value < 2958466:
            return (datetime(1899, 12, 30, tzinfo=tz) + timedelta(days=float(value))).astimezone(timezone.utc)
        return None
    s = str(value).strip()
    if re.fullmatch(r"\d{8}", s):
        try:
            return datetime(int(s[:4]), int(s[4:6]), int(s[6:]), tzinfo=tz).astimezone(timezone.utc)
        except ValueError:
            return None
    if re.fullmatch(r"\d+(\.\d+)?", s):
        return parse_date(float(s), tz)
    if re.match(r"^\d{4}-\d{1,2}-\d{1,2}", s):
        try:
            d = datetime.fromisoformat(s.replace(" ", "T", 1) if "T" not in s and " " in s else s)
        except ValueError:
            return None
        return (d.replace(tzinfo=tz) if d.tzinfo is None else d).astimezone(timezone.utc)
    m = _DMY.match(s)
    if m:
        a, b, year = int(m.group(1)), int(m.group(3)), int(m.group(4))
        hh, mm, ss = int(m.group(5) or 0), int(m.group(6) or 0), int(m.group(7) or 0)
        ampm = (m.group(8) or "").upper()
        if ampm:
            month, day = a, b
            hh = (hh % 12) + (12 if ampm == "PM" else 0)
        else:
            day, month = a, b
        try:
            return datetime(year, month, day, hh, mm, ss, tzinfo=tz).astimezone(timezone.utc)
        except ValueError:
            return None
    return None


def source_for(filename: str, settings: Settings) -> str | None:
    base = filename.lower()
    best = None
    for label, prefix in settings.import_sources.items():
        p = prefix.lower()
        if p and base.startswith(p) and (best is None or len(p) > len(settings.import_sources[best])):
            best = label
    return best


# ------------------------------------------------------------------ lista de APIs a partir de ficheiro
API_COLUMNS = {
    "name": ["displaylabel", "nome", "name", "label"],
    "application": ["aplicacao", "application", "app", "aplicacoes"],
    "expiry": ["datadeexpiracao", "dataexpiracao", "datadevalidade", "expirationdate", "expirydate", "enddate",
               "validade", "validto", "expira"],
    "expired": ["expirado", "expired", "estaexpirado"],
    "location": ["localizacao", "location", "local"],
    "registered": ["dataderegistonacmdb", "dataderegisto", "createtime", "datacriacao", "datadecriacao", "registo"],
    "url": ["url", "endpoint", "endpointurl", "endereco", "enderecourl"],
    "ucmdb_id": ["ucmdbid", "globalid", "id", "cmdbid"],
}
_BOOL = {"true": True, "sim": True, "yes": True, "1": True, "false": False, "nao": False, "não": False, "no": False, "0": False}


def find_api_file(settings: Settings) -> Path | None:
    """Ficheiro mais recente com o prefixo CERTWATCH_APIS_FILE (ex.: api-cmdb.csv) na diretoria de importação."""
    d, prefix = settings.import_dir, settings.apis_file
    if not prefix or not d.is_dir():
        return None
    files = [p for p in d.iterdir() if p.is_file() and p.name.lower().startswith(prefix)
             and p.suffix.lower() in ALLOWED_EXT]
    return max(files, key=lambda p: p.stat().st_mtime) if files else None


def read_api_file(path: Path, settings: Settings) -> tuple[list[dict], dict]:
    """Lê a exportação da lista de APIs da CMDB (CSV/TSV/Excel, cabeçalhos "[API secrets] : Display Label"...)."""
    from .fontes.ucmdb import parse_api_url
    tz = ZoneInfo(settings.timezone)
    headers, rows = read_table(path)
    normed = [norm_header(h) for h in headers]
    cols = {}
    for field, cands in API_COLUMNS.items():
        for cand in cands:
            if cand in normed:
                cols[field] = normed.index(cand)
                break
    if "name" not in cols:
        raise ValueError(f"{path.name}: coluna de nome (Display Label) não encontrada. Cabeçalhos: {', '.join(headers)[:300]}")
    out, seen, problems = [], defaultdict(int), []
    for n, row in enumerate(rows, start=2):
        cell = lambda f: (row[cols[f]] if f in cols and cols[f] < len(row) else None)
        name = _fix_mojibake(str(cell("name") or "")).strip()
        if not name:
            continue
        raw_exp = cell("expiry")
        expiry = parse_date(raw_exp, tz)
        if raw_exp not in (None, "") and expiry is None and len(problems) < MAX_PROBLEMS:
            problems.append(f"Linha {n}: data de expiração não reconhecida ({raw_exp!r}) para {name}")
        app = _fix_mojibake(str(cell("application") or "")).strip()
        loc = _fix_mojibake(str(cell("location") or "")).strip()
        raw_url = str(cell("url") or "").strip()
        scheme, host, port = parse_api_url(raw_url)
        uid = str(cell("ucmdb_id") or "").strip()
        if not uid:   # sem ID da CMDB: identificador estável a partir do nome, aplicação e localização
            base = hashlib.sha1(f"{name.lower()}|{app.lower()}|{loc.lower()}".encode()).hexdigest()[:24]
            seen[base] += 1
            uid = f"ficheiro:{base}" + (f"-{seen[base]}" if seen[base] > 1 else "")
        exp_flag = cell("expired")
        out.append({
            "ucmdb_id": uid[:64], "name": name[:512], "application": app[:512], "location": loc[:512],
            "cmdb_expired": exp_flag if isinstance(exp_flag, bool) else _BOOL.get(str(exp_flag or "").strip().lower()),
            "registered_at": parse_date(cell("registered"), tz),
            "url": raw_url[:1024] if host else "", "scheme": scheme, "host": host, "port": port,
            "expiry": expiry, "extra": {},
        })
    fields = {"nome": headers[cols["name"]], **{k: (headers[cols[f]] if f in cols else None) for k, f in (
        ("aplicacao", "application"), ("validade", "expiry"), ("expirado", "expired"), ("localizacao", "location"),
        ("registo", "registered"), ("url", "url"))}, "metodo": "ficheiro", "ficheiro": path.name}
    if problems:
        fields["avisos"] = problems[:5]
    return out, fields


def scan_import_dir(settings: Settings) -> list[dict]:
    """O que está na diretoria de importação e o que acontece a cada ficheiro (para a página)."""
    import os
    d = settings.import_dir
    if not d.is_dir():
        return []
    out = []
    for p in sorted(d.iterdir(), key=lambda x: x.name.lower()):
        if not p.is_file() or p.name.startswith(".") or p.name.lower() == "readme.md":
            continue
        name = p.name.lower()
        readable = os.access(p, os.R_OK)
        if p.suffix.lower() not in ALLOWED_EXT:
            what, ok = f"ignorado: extensão {p.suffix or '(nenhuma)'} não suportada (usar .csv, .txt ou .xlsx)", False
        elif settings.apis_file and name.startswith(settings.apis_file):
            what, ok = "lista de APIs (usada na página APIs)", True
        elif source_for(p.name, settings) is None:
            what, ok = ("ignorado: o nome não começa por nenhum prefixo conhecido ("
                        + ", ".join(f"{v}*" for v in settings.import_sources.values()) + ")"), False
        else:
            what, ok = f"fonte {source_for(p.name, settings)}", True
        if not readable:
            what, ok = f"sem permissão de leitura: correr chmod 644 import/{p.name}", False
        out.append({"name": p.name, "size": p.stat().st_size, "what": what, "ok": ok})
    return out
